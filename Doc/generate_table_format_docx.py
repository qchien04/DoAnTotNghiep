# -*- coding: utf-8 -*-
"""
Generate DOCX document with exact 2-column table format as requested by the user:
- Heading above table: "a) Chức năng [Tên chức năng]"
- Table with 2 columns:
  - Row 1: Use case       | [Tên use case]
  - Row 2: Actor          | [Tên actor]
  - Row 3: Tiền điều kiện | [Tiền điều kiện]
  - Row 4: Hậu điều kiện  | [Hậu điều kiện]
  - Row 5: Kịch bản chính | [Từng bước 1, 2, 3... kèm bảng dữ liệu mẫu]
  - Row 6: Ngoại lệ       | [Các bước ngoại lệ]
- Clean academic styling: strictly black text on white background, no colors.
"""

import os
import sys

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from data_landlord import USECASES_LANDLORD
from data_tenant import USECASES_TENANT
from data_admin import USECASES_ADMIN

DOC_DIR = r"E:\Đồ án tốt nghiệp\Doc"
OUTPUT_DOCX = os.path.join(DOC_DIR, "Tai_lieu_Dac_ta_Use_Case_He_thong_Quan_ly_Nha_tro.docx")
OUTPUT_TXT = os.path.join(DOC_DIR, "nghiệp vụ mẫu.txt")

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(tblBorders)

def add_nested_table(cell, headers, rows_data, col_widths=None):
    """Add a clean sample data table inside a table cell for steps with tabular data."""
    p_before = cell.add_paragraph()
    p_before.paragraph_format.space_before = Pt(2)
    p_before.paragraph_format.space_after = Pt(2)

    sub_table = cell.add_table(rows=len(rows_data) + 1, cols=len(headers))
    sub_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sub_table)

    # Header row
    hdr_cells = sub_table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_margins(hdr_cells[i], top=80, bottom=80, left=100, right=100)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)

    # Data rows
    for r_idx, row_values in enumerate(rows_data):
        row_cells = sub_table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = str(val)
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=90, right=90)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if c_idx == 0 or len(str(val)) <= 6 or str(val).isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0, 0, 0)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
            row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)

    # Widths
    if col_widths:
        total_given = sum(col_widths)
        scale = 4.8 / total_given if total_given > 0 else 1.0
        for row in sub_table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w * scale)

    p_after = cell.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(3)

def render_usecase_table(doc, uc_data, prefix="a)"):
    """
    Renders 1 Use Case as a 2-column table exactly matching the screenshot:
    - Title above: prefix + " " + uc_data['chuc_nang']
    - 2-Column Table:
        Row 0: Use case
        Row 1: Actor
        Row 2: Tiền điều kiện
        Row 3: Hậu điều kiện
        Row 4: Kịch bản chính
        Row 5: Ngoại lệ
    """
    # 1. Heading above the table
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(5)
    p_title.paragraph_format.keep_with_next = True
    r_t = p_title.add_run(f"{prefix} {uc_data['chuc_nang']}")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(12)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0, 0, 0)

    # 2. Add 2-column table (6 rows, 2 cols)
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    col0_width = Inches(1.5)
    col1_width = Inches(4.97)

    labels = ["Use case", "Actor", "Tiền điều kiện", "Hậu điều kiện", "Kịch bản chính", "Ngoại lệ"]

    # Fill labels in Column 0
    for row_idx, label in enumerate(labels):
        cell_lbl = table.cell(row_idx, 0)
        cell_lbl.width = col0_width
        set_cell_margins(cell_lbl, top=100, bottom=100, left=120, right=120)
        cell_lbl.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        
        # White background
        shd0 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
        cell_lbl._tc.get_or_add_tcPr().append(shd0)

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(2)
        p_lbl.paragraph_format.space_after = Pt(2)
        p_lbl.paragraph_format.line_spacing = 1.15
        r = p_lbl.add_run(label)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0, 0, 0)
        r.font.bold = False  # As in screenshot

    # Setup content cells in Column 1
    for row_idx in range(6):
        cell_val = table.cell(row_idx, 1)
        cell_val.width = col1_width
        set_cell_margins(cell_val, top=100, bottom=100, left=140, right=140)
        cell_val.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        shd1 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
        cell_val._tc.get_or_add_tcPr().append(shd1)

    # Row 0: Use case
    c0 = table.cell(0, 1)
    p0 = c0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(2)
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.line_spacing = 1.15
    r0 = p0.add_run(uc_data['use_case'])
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(11)
    r0.font.color.rgb = RGBColor(0, 0, 0)

    # Row 1: Actor
    c1 = table.cell(1, 1)
    p1 = c1.paragraphs[0]
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15
    r1 = p1.add_run(uc_data['actor'])
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(0, 0, 0)

    # Row 2: Tiền điều kiện
    c2 = table.cell(2, 1)
    p2 = c2.paragraphs[0]
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(uc_data['tien_dieu_kien'])
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(0, 0, 0)

    # Row 3: Hậu điều kiện
    c3 = table.cell(3, 1)
    p3 = c3.paragraphs[0]
    p3.paragraph_format.space_before = Pt(2)
    p3.paragraph_format.space_after = Pt(2)
    p3.paragraph_format.line_spacing = 1.15
    r3 = p3.add_run(uc_data['hau_dieu_kien'])
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(11)
    r3.font.color.rgb = RGBColor(0, 0, 0)

    # Row 4: Kịch bản chính (Steps + optional sample tables)
    c4 = table.cell(4, 1)
    # The cell already has one paragraph
    first_p = True
    for item in uc_data['kich_ban_chinh']:
        if item[0] == "text":
            if first_p:
                p_step = c4.paragraphs[0]
                first_p = False
            else:
                p_step = c4.add_paragraph()
            p_step.paragraph_format.space_before = Pt(2)
            p_step.paragraph_format.space_after = Pt(2)
            p_step.paragraph_format.line_spacing = 1.15
            p_step.paragraph_format.left_indent = Inches(0.2)
            r_step = p_step.add_run(item[1])
            r_step.font.name = 'Times New Roman'
            r_step.font.size = Pt(11)
            r_step.font.color.rgb = RGBColor(0, 0, 0)
        elif item[0] == "table":
            headers = item[1]
            rows_data = item[2]
            col_widths = item[3] if len(item) > 3 else None
            add_nested_table(c4, headers, rows_data, col_widths)

    # Row 5: Ngoại lệ
    c5 = table.cell(5, 1)
    first_p_nl = True
    if not uc_data['ngoai_le']:
        p_nl = c5.paragraphs[0]
        p_nl.paragraph_format.space_before = Pt(2)
        p_nl.paragraph_format.space_after = Pt(2)
        r_nl = p_nl.add_run("Không có ngoại lệ")
        r_nl.font.name = 'Times New Roman'
        r_nl.font.size = Pt(11)
        r_nl.font.color.rgb = RGBColor(0, 0, 0)
    else:
        for nl in uc_data['ngoai_le']:
            if first_p_nl:
                p_nl = c5.paragraphs[0]
                first_p_nl = False
            else:
                p_nl = c5.add_paragraph()
            p_nl.paragraph_format.space_before = Pt(2)
            p_nl.paragraph_format.space_after = Pt(2)
            p_nl.paragraph_format.line_spacing = 1.15
            p_nl.paragraph_format.left_indent = Inches(0.2)
            r_nl = p_nl.add_run(nl)
            r_nl.font.name = 'Times New Roman'
            r_nl.font.size = Pt(11)
            r_nl.font.color.rgb = RGBColor(0, 0, 0)

    # Spacing after each table
    p_gap = doc.add_paragraph()
    p_gap.paragraph_format.space_before = Pt(2)
    p_gap.paragraph_format.space_after = Pt(8)

def main():
    print("Building Document with 2-Column Table Structure matching user's image...")
    doc = docx.Document()

    # Page Margins
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(0.8)

    # Main Document Header
    p_uni = doc.add_paragraph()
    p_uni.paragraph_format.space_before = Pt(0)
    p_uni.paragraph_format.space_after = Pt(4)
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_uni = p_uni.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC BÁCH KHOA HÀ NỘI\n")
    r_uni.font.name = 'Times New Roman'
    r_uni.font.size = Pt(11)
    r_uni.font.bold = True
    r_uni.font.color.rgb = RGBColor(0, 0, 0)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("TÀI LIỆU ĐẶC TẢ USE CASE CHI TIẾT\nHỆ THỐNG QUẢN LÝ NHÀ TRỌ HỖ TRỢ GHÉP NGƯỜI Ở CHUNG")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(15)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(14)
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = p_sub.add_run("(Tài liệu chuẩn hóa dưới dạng bảng đặc tả Use Case chi tiết theo từng Role)")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0, 0, 0)

    # -------------------------------------------------------------
    # PHẦN 1: GIỚI THIỆU TỔNG QUAN HỆ THỐNG & CÁC TÁC NHÂN
    # -------------------------------------------------------------
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. GIỚI THIỆU TỔNG QUAN HỆ THỐNG & DANH MỤC CÁC TÁC NHÂN")
    r_h1.font.name = 'Times New Roman'
    r_h1.font.size = Pt(13.5)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0, 0, 0)

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(2)
    p_intro.paragraph_format.space_after = Pt(6)
    p_intro.paragraph_format.line_spacing = 1.2
    r_intro = p_intro.add_run(
        "Hệ thống 'Quản lý nhà trọ hỗ trợ ghép người ở chung' kết hợp đầy đủ các tính năng quản trị vận hành nhà trọ "
        "(dành cho Chủ trọ) và các tính năng tương tác, kết nối ở ghép thông minh (dành cho Người thuê). "
        "Tính năng nổi bật cốt lõi bao gồm hỗ trợ người đã có phòng hoặc chưa có phòng đăng tin tìm bạn cùng phòng, "
        "định vị bán kính trên bản đồ số, khảo sát và so khớp mức độ tương thích lối sống sinh hoạt, đồng thời hỗ trợ "
        "chủ trọ quản lý liên kết tài khoản khách thuê, chốt điện nước, tính hóa đơn và xử lý khiếu nại minh bạch."
    )
    r_intro.font.name = 'Times New Roman'
    r_intro.font.size = Pt(11)
    r_intro.font.color.rgb = RGBColor(0, 0, 0)

    # Bảng danh mục Actor
    p_act = doc.add_paragraph()
    p_act.paragraph_format.space_before = Pt(6)
    p_act.paragraph_format.space_after = Pt(3)
    r_act = p_act.add_run("Bảng 1.1: Danh mục các tác nhân (Actors) tham gia vào hệ thống")
    r_act.font.name = 'Times New Roman'
    r_act.font.size = Pt(11)
    r_act.font.bold = True
    r_act.font.color.rgb = RGBColor(0, 0, 0)

    t_actor = doc.add_table(rows=4, cols=3)
    t_actor.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_actor)
    actor_headers = ["STT", "Tác nhân (Actor)", "Mô tả vai trò và trách nhiệm"]
    actor_data = [
        ["1", "Chủ trọ (Landlord)", "Quản lý tòa nhà, phòng trọ, dịch vụ, lập hợp đồng, chốt số điện nước, phát hành hóa đơn và xử lý sự cố báo hỏng."],
        ["2", "Người thuê (Tenant)", "Tìm kiếm phòng trọ/ở ghép, đăng bài tìm bạn ở ghép (có phòng hoặc chưa có phòng), ứng tuyển nhóm, xem hóa đơn và gửi phản ánh."],
        ["3", "Quản trị viên (Admin)", "Quản trị toàn hệ thống, quản lý tài khoản người dùng, duyệt khiếu nại tranh chấp toàn sàn và cấu hình Master Data."]
    ]
    for i, h in enumerate(actor_headers):
        cell = t_actor.cell(0, i)
        cell.text = h
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
        cell._tc.get_or_add_tcPr().append(shd)

    for r_idx, row in enumerate(actor_data):
        for c_idx, val in enumerate(row):
            cell = t_actor.cell(r_idx + 1, c_idx)
            cell.text = val
            set_cell_margins(cell, 60, 60, 90, 90)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0, 0, 0)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
            cell._tc.get_or_add_tcPr().append(shd)

    # -------------------------------------------------------------
    # PHẦN 2: CÁC USE CASE DÀNH CHO CHỦ TRỌ
    # -------------------------------------------------------------
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. ĐẶC TẢ CÁC USE CASE CỦA ROLE: CHỦ TRỌ (LANDLORD)")
    r_h2.font.name = 'Times New Roman'
    r_h2.font.size = Pt(13.5)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(0, 0, 0)

    letters = ["a)", "b)", "c)", "d)", "e)", "f)", "g)", "h)", "i)", "j)", "k)", "l)"]
    for idx, uc in enumerate(USECASES_LANDLORD):
        prefix = letters[idx] if idx < len(letters) else f"{idx+1})"
        render_usecase_table(doc, uc, prefix=prefix)

    # -------------------------------------------------------------
    # PHẦN 3: CÁC USE CASE DÀNH CHO NGƯỜI THUÊ
    # -------------------------------------------------------------
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(18)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("3. ĐẶC TẢ CÁC USE CASE CỦA ROLE: NGƯỜI THUÊ (TENANT)")
    r_h3.font.name = 'Times New Roman'
    r_h3.font.size = Pt(13.5)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(0, 0, 0)

    for idx, uc in enumerate(USECASES_TENANT):
        prefix = letters[idx] if idx < len(letters) else f"{idx+1})"
        render_usecase_table(doc, uc, prefix=prefix)

    # -------------------------------------------------------------
    # PHẦN 4: CÁC USE CASE DÀNH CHO QUẢN TRỊ VIÊN (ADMIN)
    # -------------------------------------------------------------
    h4 = doc.add_heading(level=1)
    h4.paragraph_format.space_before = Pt(18)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run("4. ĐẶC TẢ CÁC USE CASE CỦA ROLE: QUẢN TRỊ VIÊN (ADMIN)")
    r_h4.font.name = 'Times New Roman'
    r_h4.font.size = Pt(13.5)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGBColor(0, 0, 0)

    for idx, uc in enumerate(USECASES_ADMIN):
        prefix = letters[idx] if idx < len(letters) else f"{idx+1})"
        render_usecase_table(doc, uc, prefix=prefix)

    doc.save(OUTPUT_DOCX)
    print(f"Done! Successfully generated {OUTPUT_DOCX} with exact table layout.")
    print(f"File size: {os.path.getsize(OUTPUT_DOCX)} bytes")

if __name__ == "__main__":
    main()
