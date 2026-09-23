# -*- coding: utf-8 -*-
"""
Script to build the final, comprehensive specification document:
- 57 separated, individual Use Cases (every CRUD operation is a standalone Use Case).
- Each Use Case title is explicitly set as Heading 2 (e.g. 'Chức năng thêm phòng trọ mới').
- Each Use Case is structured into a 2-column, 6-row table matching the user's screenshot.
- Clean academic formatting: black text, white background, no colors.
"""

import os
import sys

if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except Exception:
        pass

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

from data_split_landlord import USECASES_LANDLORD_SPLIT
from data_split_tenant import USECASES_TENANT_SPLIT
from data_split_admin import USECASES_ADMIN_SPLIT

DOC_DIR = r"E:\Đồ án tốt nghiệp\Doc"
OUTPUT_DOCX = os.path.join(DOC_DIR, "Tai_lieu_Dac_ta_Use_Case_He_thong_Quan_ly_Nha_tro.docx")
OUTPUT_TXT = os.path.join(DOC_DIR, "nghiệp vụ mẫu.txt")

def set_cell_margins(cell, top=90, bottom=90, left=130, right=130):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        r'<w:tblBorders %s>'
        r'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(tblBorders)

def add_nested_table(cell, headers, rows_data, col_widths=None):
    p_before = cell.add_paragraph()
    p_before.paragraph_format.space_before = Pt(2)
    p_before.paragraph_format.space_after = Pt(2)

    sub_table = cell.add_table(rows=len(rows_data) + 1, cols=len(headers))
    sub_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(sub_table)

    # Header
    hdr_cells = sub_table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_margins(hdr_cells[i], top=70, bottom=70, left=90, right=90)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0, 0, 0)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)

    # Data
    for r_idx, row_values in enumerate(rows_data):
        row_cells = sub_table.rows[r_idx + 1].cells
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = str(val)
            set_cell_margins(row_cells[c_idx], top=50, bottom=50, left=80, right=80)
            p = row_cells[c_idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            if c_idx == 0 or len(str(val)) <= 6 or str(val).isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(8.5)
                r.font.color.rgb = RGBColor(0, 0, 0)
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
            row_cells[c_idx]._tc.get_or_add_tcPr().append(shd)

    if col_widths:
        total_w = sum(col_widths)
        scale = 4.8 / total_w if total_w > 0 else 1.0
        for row in sub_table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w * scale)

    p_after = cell.add_paragraph()
    p_after.paragraph_format.space_before = Pt(2)
    p_after.paragraph_format.space_after = Pt(2)

def add_usecase_with_heading2(doc, uc_data):
    """
    Creates:
    1. Heading 2 for the Use Case title (e.g. 'Chức năng thêm phòng trọ mới')
    2. Exactly 2-column, 6-row table matching user's image.
    """
    # 1. HEADING 2
    h2 = doc.add_heading(level=2)
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True
    
    run_h2 = h2.add_run(f"{uc_data['stt']}. {uc_data['chuc_nang']}")
    run_h2.font.name = 'Times New Roman'
    run_h2.font.size = Pt(12.5)
    run_h2.font.bold = True
    run_h2.font.color.rgb = RGBColor(0, 0, 0)

    # 2. 2-COLUMN TABLE (6 rows, 2 cols)
    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)

    col0_width = Inches(1.5)
    col1_width = Inches(4.97)

    labels = ["Use case", "Actor", "Tiền điều kiện", "Hậu điều kiện", "Kịch bản chính", "Ngoại lệ"]

    # Fill labels in Column 0
    for row_idx, label in enumerate(labels):
        cell_lbl = table.cell(row_idx, 0)
        cell_lbl.width = col0_width
        set_cell_margins(cell_lbl, top=90, bottom=90, left=110, right=110)
        cell_lbl.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        
        shd0 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
        cell_lbl._tc.get_or_add_tcPr().append(shd0)

        p_lbl = cell_lbl.paragraphs[0]
        p_lbl.paragraph_format.space_before = Pt(2)
        p_lbl.paragraph_format.space_after = Pt(2)
        p_lbl.paragraph_format.line_spacing = 1.15
        r = p_lbl.add_run(label)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0, 0, 0)
        r.font.bold = False

    # Setup content cells in Column 1
    for row_idx in range(6):
        cell_val = table.cell(row_idx, 1)
        cell_val.width = col1_width
        set_cell_margins(cell_val, top=90, bottom=90, left=130, right=130)
        cell_val.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        shd1 = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FFFFFF"/>')
        cell_val._tc.get_or_add_tcPr().append(shd1)

    # Row 0: Use case
    p0 = table.cell(0, 1).paragraphs[0]
    p0.paragraph_format.space_before = Pt(2)
    p0.paragraph_format.space_after = Pt(2)
    p0.paragraph_format.line_spacing = 1.15
    r0 = p0.add_run(uc_data['use_case'])
    r0.font.name = 'Times New Roman'
    r0.font.size = Pt(11)
    r0.font.color.rgb = RGBColor(0, 0, 0)

    # Row 1: Actor
    p1 = table.cell(1, 1).paragraphs[0]
    p1.paragraph_format.space_before = Pt(2)
    p1.paragraph_format.space_after = Pt(2)
    p1.paragraph_format.line_spacing = 1.15
    r1 = p1.add_run(uc_data['actor'])
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.font.color.rgb = RGBColor(0, 0, 0)

    # Row 2: Tiền điều kiện
    p2 = table.cell(2, 1).paragraphs[0]
    p2.paragraph_format.space_before = Pt(2)
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.15
    r2 = p2.add_run(uc_data['tien_dieu_kien'])
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    r2.font.color.rgb = RGBColor(0, 0, 0)

    # Row 3: Hậu điều kiện
    p3 = table.cell(3, 1).paragraphs[0]
    p3.paragraph_format.space_before = Pt(2)
    p3.paragraph_format.space_after = Pt(2)
    p3.paragraph_format.line_spacing = 1.15
    r3 = p3.add_run(uc_data['hau_dieu_kien'])
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(11)
    r3.font.color.rgb = RGBColor(0, 0, 0)

    # Row 4: Kịch bản chính
    c4 = table.cell(4, 1)
    first_p = True
    for item in uc_data['kich_ban_chinh']:
        if item[0] == "text":
            if first_p:
                p_step = c4.paragraphs[0]
                first_p = False
            else:
                p_step = c4.add_paragraph()
            p_step.paragraph_format.space_before = Pt(2)
            p_step.paragraph_format.space_after = Pt(2)
            p_step.paragraph_format.line_spacing = 1.15
            p_step.paragraph_format.left_indent = Inches(0.2)
            r_step = p_step.add_run(item[1])
            r_step.font.name = 'Times New Roman'
            r_step.font.size = Pt(11)
            r_step.font.color.rgb = RGBColor(0, 0, 0)
        elif item[0] == "table":
            headers = item[1]
            rows_data = item[2]
            col_widths = item[3] if len(item) > 3 else None
            add_nested_table(c4, headers, rows_data, col_widths)

    # Row 5: Ngoại lệ
    c5 = table.cell(5, 1)
    if not uc_data['ngoai_le']:
        p_nl = c5.paragraphs[0]
        p_nl.paragraph_format.space_before = Pt(2)
        p_nl.paragraph_format.space_after = Pt(2)
        r_nl = p_nl.add_run("Không có ngoại lệ")
        r_nl.font.name = 'Times New Roman'
        r_nl.font.size = Pt(11)
        r_nl.font.color.rgb = RGBColor(0, 0, 0)
    else:
        first_p_nl = True
        for nl in uc_data['ngoai_le']:
            if first_p_nl:
                p_nl = c5.paragraphs[0]
                first_p_nl = False
            else:
                p_nl = c5.add_paragraph()
            p_nl.paragraph_format.space_before = Pt(2)
            p_nl.paragraph_format.space_after = Pt(2)
            p_nl.paragraph_format.line_spacing = 1.15
            p_nl.paragraph_format.left_indent = Inches(0.2)
            r_nl = p_nl.add_run(nl)
            r_nl.font.name = 'Times New Roman'
            r_nl.font.size = Pt(11)
            r_nl.font.color.rgb = RGBColor(0, 0, 0)

    # Spacing after table
    p_gap = doc.add_paragraph()
    p_gap.paragraph_format.space_before = Pt(2)
    p_gap.paragraph_format.space_after = Pt(8)

def export_text_version(all_usecases):
    lines = []
    lines.append("=" * 80)
    lines.append("HỆ THỐNG QUẢN LÝ NHÀ TRỌ HỖ TRỢ GHÉP NGƯỜI Ở CHUNG")
    lines.append("TÀI LIỆU ĐẶC TẢ CHI TIẾT 57 USE CASE (TÁCH BIỆT TỪNG THAO TÁC CRUD)")
    lines.append("=" * 80 + "\n")

    current_role = None
    for uc in all_usecases:
        # Determine role from index
        if uc['stt'] == 1:
            lines.append("\n" + "#" * 80)
            lines.append("PHẦN 1: CÁC USE CASE DÀNH CHO ROLE: CHỦ TRỌ (LANDLORD)")
            lines.append("#" * 80 + "\n")
        elif uc['stt'] == 30:
            lines.append("\n" + "#" * 80)
            lines.append("PHẦN 2: CÁC USE CASE DÀNH CHO ROLE: NGƯỜI THUÊ (TENANT)")
            lines.append("#" * 80 + "\n")
        elif uc['stt'] == 47:
            lines.append("\n" + "#" * 80)
            lines.append("PHẦN 3: CÁC USE CASE DÀNH CHO ROLE: QUẢN TRỊ VIÊN (ADMIN)")
            lines.append("#" * 80 + "\n")

        lines.append(f"Header 2: {uc['stt']}. {uc['chuc_nang']}")
        lines.append("-" * 60)
        lines.append(f"Use case: {uc['use_case']}")
        lines.append(f"Actor: {uc['actor']}")
        lines.append(f"Tiền điều kiện: {uc['tien_dieu_kien']}")
        lines.append(f"Hậu điều kiện: {uc['hau_dieu_kien']}")
        lines.append("Kịch bản chính:")
        for item in uc['kich_ban_chinh']:
            if item[0] == "text":
                lines.append(f"  {item[1]}")
            elif item[0] == "table":
                headers = item[1]
                rows = item[2]
                lines.append("  " + " | ".join(headers))
                lines.append("  " + " | ".join(["---"] * len(headers)))
                for r in rows:
                    lines.append("  " + " | ".join(r))
        lines.append("Ngoại lệ:")
        if not uc['ngoai_le']:
            lines.append("  Không có ngoại lệ")
        else:
            for nl in uc['ngoai_le']:
                lines.append(f"  {nl}")
        lines.append("\n" + "=" * 60 + "\n")

    return "\n".join(lines)

def main():
    print("Starting generation of separated CRUD use cases document...")
    doc = docx.Document()

    # Configure styles in Word document
    styles = doc.styles
    
    # Customise Heading 1
    h1_style = styles['Heading 1']
    h1_style.font.name = 'Times New Roman'
    h1_style.font.size = Pt(14)
    h1_style.font.bold = True
    h1_style.font.color.rgb = RGBColor(0, 0, 0)

    # Customise Heading 2
    h2_style = styles['Heading 2']
    h2_style.font.name = 'Times New Roman'
    h2_style.font.size = Pt(12.5)
    h2_style.font.bold = True
    h2_style.font.color.rgb = RGBColor(0, 0, 0)

    # Normal Style
    normal_style = styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    # Page Margins
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(0.8)

    # Document Title
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_uni.paragraph_format.space_before = Pt(0)
    p_uni.paragraph_format.space_after = Pt(4)
    r_uni = p_uni.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO\nTRƯỜNG ĐẠI HỌC BÁCH KHOA HÀ NỘI\n")
    r_uni.font.name = 'Times New Roman'
    r_uni.font.size = Pt(11)
    r_uni.font.bold = True
    r_uni.font.color.rgb = RGBColor(0, 0, 0)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(10)
    p_title.paragraph_format.space_after = Pt(4)
    r_t = p_title.add_run("TÀI LIỆU ĐẶC TẢ CHI TIẾT CÁC USE CASE\nHỆ THỐNG QUẢN LÝ NHÀ TRỌ HỖ TRỢ GHÉP NGƯỜI Ở CHUNG")
    r_t.font.name = 'Times New Roman'
    r_t.font.size = Pt(15)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0, 0, 0)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(2)
    p_sub.paragraph_format.space_after = Pt(16)
    r_sub = p_sub.add_run("(Tách biệt hoàn toàn từng chức năng Thêm - Sửa - Xóa - Xem; Đánh Header 2 cho từng Use Case)")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(11)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(0, 0, 0)

    # Intro overview
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(4)
    p_intro.paragraph_format.space_after = Pt(12)
    p_intro.paragraph_format.line_spacing = 1.2
    r_intro = p_intro.add_run(
        "Tài liệu này đặc tả chi tiết toàn bộ 57 Use Case của hệ thống 'Quản lý nhà trọ hỗ trợ ghép người ở chung'. "
        "Mỗi thao tác nghiệp vụ (Thêm, Sửa, Xóa, Xem danh sách, Duyệt, Khóa, Thanh lý, Gửi thông báo...) được xây dựng thành một Use Case độc lập "
        "và được phân cấp định dạng Header 2 để phục vụ tạo mục lục tự động và đánh giá kiến trúc phần mềm chuẩn xác. "
        "Mỗi Use Case được trình bày dạng bảng 2 cột theo đúng chuẩn đặc tả ca sử dụng."
    )
    r_intro.font.name = 'Times New Roman'
    r_intro.font.size = Pt(11)
    r_intro.font.color.rgb = RGBColor(0, 0, 0)

    # -------------------------------------------------------------
    # PHẦN 1: CHỦ TRỌ (Landlord) - 29 Use Cases
    # -------------------------------------------------------------
    h1_ll = doc.add_heading(level=1)
    h1_ll.paragraph_format.space_before = Pt(18)
    h1_ll.paragraph_format.space_after = Pt(8)
    r_h1_ll = h1_ll.add_run("1. CÁC USE CASE DÀNH CHO ROLE: CHỦ TRỌ (LANDLORD)")
    r_h1_ll.font.name = 'Times New Roman'
    r_h1_ll.font.size = Pt(13.5)
    r_h1_ll.font.bold = True
    r_h1_ll.font.color.rgb = RGBColor(0, 0, 0)

    for uc in USECASES_LANDLORD_SPLIT:
        add_usecase_with_heading2(doc, uc)

    # -------------------------------------------------------------
    # PHẦN 2: NGƯỜI THUÊ (Tenant) - 17 Use Cases
    # -------------------------------------------------------------
    h1_tn = doc.add_heading(level=1)
    h1_tn.paragraph_format.space_before = Pt(18)
    h1_tn.paragraph_format.space_after = Pt(8)
    r_h1_tn = h1_tn.add_run("2. CÁC USE CASE DÀNH CHO ROLE: NGƯỜI THUÊ (TENANT)")
    r_h1_tn.font.name = 'Times New Roman'
    r_h1_tn.font.size = Pt(13.5)
    r_h1_tn.font.bold = True
    r_h1_tn.font.color.rgb = RGBColor(0, 0, 0)

    for uc in USECASES_TENANT_SPLIT:
        add_usecase_with_heading2(doc, uc)

    # -------------------------------------------------------------
    # PHẦN 3: QUẢN TRỊ VIÊN (Admin) - 11 Use Cases
    # -------------------------------------------------------------
    h1_ad = doc.add_heading(level=1)
    h1_ad.paragraph_format.space_before = Pt(18)
    h1_ad.paragraph_format.space_after = Pt(8)
    r_h1_ad = h1_ad.add_run("3. CÁC USE CASE DÀNH CHO ROLE: QUẢN TRỊ VIÊN (ADMIN)")
    r_h1_ad.font.name = 'Times New Roman'
    r_h1_ad.font.size = Pt(13.5)
    r_h1_ad.font.bold = True
    r_h1_ad.font.color.rgb = RGBColor(0, 0, 0)

    for uc in USECASES_ADMIN_SPLIT:
        add_usecase_with_heading2(doc, uc)

    # Save Word document
    doc.save(OUTPUT_DOCX)
    print(f"Successfully generated DOCX with {len(USECASES_LANDLORD_SPLIT) + len(USECASES_TENANT_SPLIT) + len(USECASES_ADMIN_SPLIT)} Use Cases!")
    print(f"DOCX Location: {OUTPUT_DOCX}")
    print(f"DOCX File Size: {os.path.getsize(OUTPUT_DOCX)} bytes")

    # Export Text version
    all_ucs = USECASES_LANDLORD_SPLIT + USECASES_TENANT_SPLIT + USECASES_ADMIN_SPLIT
    txt_content = export_text_version(all_ucs)
    with open(OUTPUT_TXT, "w", encoding="utf-8") as f:
        f.write(txt_content)
    print(f"Text version updated at: {OUTPUT_TXT}")
    print(f"Text File Size: {os.path.getsize(OUTPUT_TXT)} bytes")

if __name__ == "__main__":
    main()
